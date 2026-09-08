"""Eurostat current EA21 all-items HICP annual rate, including labelled estimates.

2026 uses prc_hicp_minr/coicop18=TOTAL. prc_hicp_manr ended in 2025.
https://ec.europa.eu/eurostat/cache/metadata/en/prc_hicp_esms.htm
"""
from datetime import datetime, timezone
import math
import requests
from official_macro import EUROSTAT_BASE, _category_index, _period_end, _utc_now


def parse_hicp(payload, *, now=None, geo="EA21"):
    if geo not in ("EA21", "CH"):
        raise ValueError("HICP_GEOGRAPHY_NOT_APPROVED")
    checked = _utc_now(now)
    if not isinstance(payload, dict) or payload.get('class') != 'dataset' or payload.get('source') != 'ESTAT' or payload.get('extension', {}).get('id', '').lower() != 'prc_hicp_minr':
        raise ValueError('HICP_DATASET_INVALID')
    expected = {'freq': 'M', 'unit': 'RCH_A', 'coicop18': 'TOTAL', 'geo': geo}
    dimensions, ids, sizes = payload.get('dimension', {}), payload.get('id', []), payload.get('size', [])
    if len(ids) != len(set(ids)) or set(ids) != set(expected) | {'time'} or len(sizes) != len(ids):
        raise ValueError('HICP_DIMENSIONS_INVALID')
    for key, value in expected.items():
        if _category_index(dimensions[key]['category']) != {value: 0} or sizes[ids.index(key)] != 1:
            raise ValueError('HICP_SERIES_MISMATCH')
    times = _category_index(dimensions['time']['category'])
    if sizes[ids.index('time')] != len(times):
        raise ValueError('HICP_TIME_SHAPE_INVALID')
    values, statuses = payload.get('value', {}), payload.get('status', {})
    if isinstance(values, list):
        if len(values) != len(times):
            raise ValueError('HICP_VALUE_SHAPE_INVALID')
        values = {str(i): value for i, value in enumerate(values)}
    if not isinstance(values, dict) or any(key not in {str(i) for i in range(len(times))} for key in values):
        raise ValueError('HICP_VALUES_INVALID')
    if not isinstance(statuses, (dict, list)):
        raise ValueError('HICP_STATUS_INVALID')
    observations = []
    for period, position in times.items():
        date = _period_end(period, 'M')
        value = values.get(str(position))
        if value is None:
            continue
        if type(value) not in (float, int) or not math.isfinite(value) or not -25 <= value <= 25 or date > checked.date():
            raise ValueError('HICP_RATE_OR_PERIOD_INVALID')
        status = statuses.get(str(position)) if isinstance(statuses, dict) else (statuses[position] if position < len(statuses) else None)
        flags = set(str(status or '').split())
        if flags - {'e', 'p', 'b', 'd', 'r'}:
            raise ValueError('HICP_STATUS_NOT_APPROVED')
        observations.append((date, period, float(value), status, 'e' in flags))
    if not observations:
        return None
    date, period, value, status, estimate = max(observations)
    return {'value': value, 'date': date.isoformat(), 'reference_period': period,
            'source': 'Eurostat', 'source_url': EUROSTAT_BASE + 'prc_hicp_minr',
            'series_id': 'prc_hicp_minr:M.RCH_A.TOTAL.' + geo, 'frequency': 'monthly',
            'unit': 'annual percent change', 'seasonal_adjustment': 'NSA',
            'checked_at': checked.isoformat(), 'published_at': None,
            'provider_status': status, 'is_estimate': estimate}


def fetch_hicp(*, now=None, session=None, geo="EA21"):
    if geo not in ("EA21", "CH"):
        raise ValueError("HICP_GEOGRAPHY_NOT_APPROVED")
    checked = now or datetime.now(timezone.utc)
    response = (session or requests).get(EUROSTAT_BASE + 'prc_hicp_minr',
        params={'freq': 'M', 'unit': 'RCH_A', 'coicop18': 'TOTAL', 'geo': geo,
                'sinceTimePeriod': f'{checked.year - 1}-01', 'lang': 'EN'}, timeout=20)
    response.raise_for_status()
    return parse_hicp(response.json(), now=checked, geo=geo)


# Public FSO page component supplies the current workbook ID each month.
BFS_HICP_PAGE = 'https://www.bfs.admin.ch/bfs/de/home/statistiken/preise/harmonisierte-verbraucherpreise.html'
BFS_HICP_DISCOVERY = ('https://www.bfs.admin.ch/content/bfs/de/home/statistiken/preise/'
    'harmonisierte-verbraucherpreise/jcr:content/root/main/section/container/'
    'tabs_419989973/item_1/ws_composed_list.model.json')
BFS_ASSET_BASE = 'https://dam-api.bfs.admin.ch/hub/api/dam/assets/'


def discover_swiss_hicp(payload):
    """Resolve the official current HICP table; ambiguity fails closed."""
    import json
    try:
        listing = json.loads(payload['assetListAsJson'])
        candidates = [x for x in listing['list']
                      if x['title'].startswith('HVPI Schweiz (')
                      and 'Detailresultate seit' in x['title']]
        if len(candidates) != 1:
            raise ValueError()
        item = candidates[0]
        if type(item['damId']) is not int or item['damId'] <= 0:
            raise ValueError()
        expected = BFS_ASSET_BASE + str(item['damId'])
        if item['url'] != expected:
            raise ValueError()
        return expected
    except (KeyError, TypeError, ValueError) as exc:
        raise ValueError('BFS_HICP_DISCOVERY_INVALID') from exc


def parse_swiss_hicp(content, metadata, *, now=None):
    """FSO all-items monthly HICP YoY, never the national CPI/LIK."""
    import io
    import re
    import openpyxl
    checked = _utc_now(now)
    try:
        description = metadata['description']
        categories = description['categorization']
        for field, code in [('termsOfUse', 'OPEN-BY'), ('dataSource', 'BFS'),
                            ('inquiry', 'HVPI'), ('periodicity', 'MONTHLY')]:
            if code not in {x.get('code') for x in categories[field]}:
                raise ValueError()
        if {x.get('code') for x in categories['termsOfUse']} != {'OPEN-BY'}:
            raise ValueError()
        source_title = description['titles']['main']
        if (not isinstance(source_title, str) or len(source_title) > 180
                or not source_title.startswith('HVPI Schweiz (')
                or 'Detailresultate seit' not in source_title
                or any(ord(char) < 32 for char in source_title)):
            raise ValueError()
        provisional = metadata['bfs']['provisional']
        if type(provisional) is not bool:
            raise ValueError()
        if metadata['bfs']['lifecycleGroup'] != 'CURRENT':
            raise ValueError()
        published = datetime.fromisoformat(metadata['bfs']['embargo'].replace('Z', '+00:00'))
        if published.tzinfo is None or published > checked:
            raise ValueError()
        reference_end = datetime.strptime(description['bibliography']['period'].split('-')[-1], '%d.%m.%Y').date()
        period = reference_end.strftime('%Y-%m')
        if _period_end(period, 'M') != reference_end or reference_end > published.date():
            raise ValueError()
        asset_id = metadata['ids']['damId']
        if type(asset_id) is not int:
            raise ValueError()
        source_url = BFS_ASSET_BASE + str(asset_id) + '/master'
        masters = [x for x in metadata['links'] if x['rel'] == 'master']
        if len(masters) != 1 or masters[0].get('format') != 'xlsx' or masters[0]['href'] != source_url:
            raise ValueError()
    except (KeyError, TypeError, ValueError, AttributeError) as exc:
        raise ValueError('BFS_HICP_METADATA_INVALID') from exc
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = workbook['%_m-12']
            # FSO files currently declare A1:A1 although the full data are present.
            sheet.reset_dimensions()
            rows = list(sheet.iter_rows(min_row=1, max_row=8, values_only=True))
            if ('HARMONISED INDEX OF CONSUMER PRICES (HICP)' not in str(rows[0][0])
                    or not str(rows[2][0]).startswith('Annual rate of change /')
                    or str(rows[4][0]).strip() != 'Switzerland / Suisse / Schweiz'
                    or list(rows[6][:5]) != ['ITEM', 'Type de position', 'EN', 'FR', 'DE']
                    or list(rows[7][:5]) != ['_T', 1, 'All items', 'Total', 'Total']):
                raise ValueError()
            dates = list(rows[6][5:])
            values = list(rows[7][5:])
            # No empty future placeholders, duplicate periods, or skipped months.
            if not dates or any(not isinstance(x, str) or not re.fullmatch(r'\d{4}-\d{2}', x) for x in dates):
                raise ValueError()
            serial = [int(x[:4]) * 12 + int(x[5:]) for x in dates]
            if any(_period_end(x, 'M') > reference_end for x in dates) or any(b != a + 1 for a, b in zip(serial, serial[1:])):
                raise ValueError()
            if dates[-1] != period or len(values) != len(dates):
                raise ValueError()
            value = values[-1]
            if type(value) not in (int, float) or not math.isfinite(value) or not -25 <= value <= 25:
                raise ValueError()
        finally:
            workbook.close()
    except Exception as exc:
        raise ValueError('BFS_HICP_WORKBOOK_INVALID') from exc
    return {'value': float(value), 'date': reference_end.isoformat(), 'reference_period': period,
            'source': 'BFS', 'source_title': source_title, 'source_url': source_url, 'series_id': 'BFS:HICP:CH:_T:M:RCH_A',
            'frequency': 'monthly', 'unit': 'annual percent change', 'seasonal_adjustment': 'NSA',
            'checked_at': checked.isoformat(), 'published_at': published.isoformat(),
            'next_due_at': None, 'is_estimate': provisional,
            'redistribution_status': 'permitted_with_attribution', 'license': 'OPEN-BY'}


def fetch_swiss_hicp(*, now=None, session=None):
    client = session or requests
    def get(url):
        response = client.get(url, timeout=30)
        response.raise_for_status()
        return response
    def get_json(url, error_code):
        response = get(url)  # Transport errors retain their RequestException type.
        try:
            payload = response.json()
        except ValueError as exc:
            # requests.JSONDecodeError also inherits RequestException; normalize
            # malformed successful responses before the app's transport handler.
            raise ValueError(error_code) from exc
        if not isinstance(payload, dict):
            raise ValueError(error_code)
        return payload
    asset_url = discover_swiss_hicp(get_json(BFS_HICP_DISCOVERY, 'BFS_HICP_DISCOVERY_INVALID'))
    metadata = get_json(asset_url, 'BFS_HICP_METADATA_INVALID')
    if BFS_ASSET_BASE + str(metadata.get('ids', {}).get('damId')) != asset_url:
        raise ValueError('BFS_HICP_ASSET_ID_MISMATCH')
    # Construct a fixed official-domain URL; never follow metadata-controlled hosts.
    content = get(asset_url + '/master').content
    return parse_swiss_hicp(content, metadata, now=now)
