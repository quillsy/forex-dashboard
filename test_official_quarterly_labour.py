import copy
import csv
import html
import io
import json
import unittest
from datetime import datetime, timezone

from official_quarterly_labour import (
    CH_RIGHTS, CH_TITLE, fetch_ch_labour, fetch_nz_labour, fetch_quarterly_labour,
    parse_ch_labour, parse_nz_labour, select_ch_labour_resources,
)

NOW = datetime(2026, 9, 7, 12, tzinfo=timezone.utc)
URL = 'https://dam-api.bfs.admin.ch/hub/api/dam/assets/123456/master'
ISSUED = '2026-08-18T06:30:00+00:00'


def nz_document():
    return {'Title': 'Labour market statistics: June 2026 quarter',
            'DateTaxonomyTerm': {'PublicationDate': '2026-08-05 10:45:00'},
            'FeaturedMedia': {'GraphHeading': 'Unemployment rate by sex, seasonally adjusted, June 2012–June 2026 quarters',
                'SeriesData': [{'GraphCsvData': 'Quarter,Sept-25,Mar-26,Jun-26\nMen,5.2,5.4,5.7\nWomen,5.1,5.3,5.5\nTotal,5.2,5.4,5.6\n'}]},
            'Blocks': [{'HTML': '<p>Labour market statistics: September 2026 quarter</p> will be released on <b>4 November 2026</b>.'}]}


def page(d):
    return '<div data-value="' + html.escape(json.dumps(d), quote=True) + '"></div>'


def ch_csv(rows=None):
    fields = ['INDICATORS_HRCHY', 'INDICATORS_DE', 'SEASON_D', 'DETAILS_DE', 'PERIOD', 'FREQ', 'MEASURE_DE', 'VALUE', 'STATUS']
    row = ['P', 'TOTAL', 'saisonbereinigte', 'Total', '2026-Q2', 'Q', 'Durchschnittliche Quartalswerte', '5.11906217648145', 'A']
    out = io.StringIO(); writer = csv.writer(out); writer.writerow(fields)
    writer.writerows(rows if rows is not None else [row])
    return out.getvalue().encode('utf-8-sig')


def catalog():
    return {'success': True, 'result': {'count': 1, 'results': [{
        'name': 'ilo-quarterly-labour', 'title': {'de': CH_TITLE}, 'organization': {'name': 'bundesamt-fur-statistik-bfs'},
        'resources': [{'format': 'CSV', 'url': URL, 'rights': CH_RIGHTS, 'issued': ISSUED}]}]}}


class Response:
    def __init__(self, text='', content=b'', payload=None, status=200):
        self.text, self.content, self.payload, self.status_code = text, content, payload, status
    def raise_for_status(self):
        if self.status_code >= 400: raise RuntimeError('HTTP_FAILED')
    def json(self): return self.payload


class Session:
    def __init__(self, responses): self.responses, self.calls = list(responses), []
    def get(self, url, **kwargs):
        self.calls.append((url, kwargs)); return self.responses.pop(0)


class QuarterlyLabourTests(unittest.TestCase):
    def test_common_entrypoint_and_currency_guard(self):
        session = Session([Response(text=page(nz_document()))])
        row = fetch_quarterly_labour("NZD", now=NOW, session=session)
        self.assertEqual(row["period_label"], "Saisonbereinigte Quartalsquote")
        with self.assertRaises(ValueError): fetch_quarterly_labour("USD", now=NOW, session=session)

    def test_nz_total_actual_publication_and_conservative_due(self):
        row = parse_nz_labour(page(nz_document()), year=2026, quarter=2, now=NOW)
        self.assertEqual(row['value'], 5.6)
        self.assertEqual(row['date'], '2026-06-30')
        self.assertEqual(row['frequency'], 'quarterly')
        self.assertEqual(row['published_at'], '2026-08-04T22:45:00+00:00')
        self.assertEqual(row['next_due_at'], '2026-11-03T11:00:00+00:00')

    def test_nz_rejects_wrong_survey_missing_total_future_and_conflict(self):
        cases = []
        d = nz_document(); d['FeaturedMedia']['GraphHeading'] = 'Underutilisation rate'; cases.append(d)
        d = nz_document(); d['FeaturedMedia']['SeriesData'][0]['GraphCsvData'] = 'Quarter,Jun-26\nMen,5.7'; cases.append(d)
        d = nz_document(); d['DateTaxonomyTerm']['PublicationDate'] = '2026-10-05 10:45:00'; cases.append(d)
        d = nz_document(); d['FeaturedMedia']['SeriesData'].append({'GraphCsvData': 'Quarter,Jun-26\nTotal,9.9'}); cases.append(d)
        d = nz_document(); d['FeaturedMedia']['SeriesData'][0]['GraphCsvData'] = 'Quarter,Jun-26,Sep-26\nTotal,5.6,5.7'; cases.append(d)
        d = nz_document(); d['FeaturedMedia']['SeriesData'][0]['GraphCsvData'] = 'Quarter,Jun-26\nTotal,nan'; cases.append(d)
        for d in cases:
            with self.subTest(d=d), self.assertRaises(ValueError):
                parse_nz_labour(page(d), year=2026, quarter=2, now=NOW)

    def test_nz_rejects_duplicate_period_and_conflicting_calendar(self):
        d = nz_document(); d['FeaturedMedia']['SeriesData'][0]['GraphCsvData'] = 'Quarter,Jun-26,Jun-26\nTotal,5.6,5.6'
        with self.assertRaises(ValueError): parse_nz_labour(page(d), year=2026, quarter=2, now=NOW)
        d = nz_document(); d['Blocks'].append('Labour market statistics: September 2026 quarter will be released on 5 November 2026')
        with self.assertRaises(ValueError): parse_nz_labour(page(d), year=2026, quarter=2, now=NOW)

    def test_nz_unknown_next_date_remains_unknown(self):
        d = nz_document(); d.pop('Blocks')
        self.assertIsNone(parse_nz_labour(page(d), year=2026, quarter=2, now=NOW)['next_due_at'])

    def test_nz_discovery_only_404_permits_previous_quarter(self):
        session = Session([Response(status=404), Response(text=page(nz_document()))])
        row = fetch_nz_labour(now=datetime(2026, 10, 7, tzinfo=timezone.utc), session=session)
        self.assertEqual(row['reference_period'], '2026-Q2')
        self.assertIn('september-2026', session.calls[0][0]); self.assertIn('june-2026', session.calls[1][0])
        session = Session([Response(status=503), Response(text=page(nz_document()))])
        with self.assertRaises(RuntimeError): fetch_nz_labour(now=NOW, session=session)
        self.assertEqual(len(session.calls), 1)

    def test_ch_selects_latest_official_published_resource(self):
        d = catalog(); resources = d['result']['results'][0]['resources']
        resources.append(dict(resources[0], url=URL.replace('123456', '999999'), issued='2026-05-18T06:30:00Z'))
        resources.append(dict(resources[0], url=URL.replace('123456', '888888'), issued='2026-11-18T06:30:00Z'))
        self.assertEqual(select_ch_labour_resources(d, now=NOW), [{'issued': ISSUED, 'url': URL, 'source_url': 'https://opendata.swiss/dataset/ilo-quarterly-labour'}])

    def test_ch_rejects_wrong_rights_origin_and_truncated_catalog(self):
        for field, value in [('rights', 'https://opendata.swiss/terms-of-use#terms_by_ask'), ('url', 'https://example.com/data.csv')]:
            d = catalog(); d['result']['results'][0]['resources'][0][field] = value
            with self.assertRaises(ValueError): select_ch_labour_resources(d, now=NOW)
        d = catalog(); d['result']['count'] = 101
        with self.assertRaises(ValueError): select_ch_labour_resources(d, now=NOW)
        d = catalog(); d['result']['results'][0]['organization']['name'] = 'other'
        with self.assertRaises(ValueError): select_ch_labour_resources(d, now=NOW)

    def test_ch_total_sa_value(self):
        row = parse_ch_labour(ch_csv(), published_at=ISSUED, source_url=URL, now=NOW)
        self.assertEqual(row['reference_period'], '2026-Q2'); self.assertAlmostEqual(row['value'], 5.11906217648145)
        self.assertEqual(row['published_at'], ISSUED); self.assertIsNone(row['next_due_at'])

    def test_ch_rejects_bad_metadata_future_nan_conflicts(self):
        original = ['P', 'TOTAL', 'saisonbereinigte', 'Total', '2026-Q2', 'Q', 'Durchschnittliche Quartalswerte', '5.1', 'A']
        for position, value in [(1, 'OTHER'), (3, 'Men'), (4, '2026-Q4'), (5, 'M'), (6, 'Monatswerte'), (7, 'nan'), (8, 'O')]:
            row = original[:]; row[position] = value
            with self.subTest(position=position), self.assertRaises(ValueError):
                parse_ch_labour(ch_csv([row]), published_at=ISSUED, source_url=URL, now=NOW)
        other = original[:]; other[7] = '5.2'
        with self.assertRaises(ValueError): parse_ch_labour(ch_csv([original, other]), published_at=ISSUED, source_url=URL, now=NOW)

    def test_ch_fetch_dynamic_url_and_conflicting_releases(self):
        session = Session([Response(payload=catalog()), Response(content=ch_csv())])
        self.assertEqual(fetch_ch_labour(now=NOW, session=session)['source_url'], 'https://opendata.swiss/dataset/ilo-quarterly-labour')
        self.assertEqual(session.calls[1][0], URL)
        d = catalog(); d['result']['results'][0]['resources'].append(dict(d['result']['results'][0]['resources'][0], url=URL.replace('123456', '777777')))
        session = Session([Response(payload=d), Response(content=ch_csv()), Response(content=ch_csv().replace(b'5.11906217648145', b'8.8'))])
        with self.assertRaises(ValueError): fetch_ch_labour(now=NOW, session=session)




class JapanLabourTests(unittest.TestCase):
    def release_pages(self):
        return ('<table><tr><td>Monthly</td><td>- July 2026 - (Released on August 28, 2026) Main results</td><td>-</td></tr></table>',
                '<table><tr><td>2026 July</td><td>August 28</td><td></td><td></td></tr>'
                '<tr><td>August</td><td>October 2</td><td></td><td></td></tr></table>')

    def book(self, mutate=None):
        from openpyxl import Workbook
        w = Workbook(); s = w.active; s.title = '季節調整値'
        s.cell(2, 5, 'Historical data 1 a-1 Major items - Whole Japan, Monthly Data')
        s.cell(5, 5, 'Seasonally adjusted series')
        s.cell(7, 20, 'Unemployment rate  (percent)'); s.cell(7, 22, '')
        s.cell(9, 20, 'Both sexes')
        for month in range(1, 13):
            row = month + 10
            if month == 1: s.cell(row, 1, '令和 8年')
            if month == 2: s.cell(row, 1, 2026)
            s.cell(row, 2, f'{month}月')
            if month <= 7: s.cell(row, 20, 2.4 if month == 7 else 2.5)
        if mutate: mutate(s)
        out = io.BytesIO(); w.save(out); return out.getvalue()

    def test_actual_layout_july_and_blank_future_months(self):
        from official_quarterly_labour import parse_japan_labour, parse_japan_labour_release
        release = parse_japan_labour_release(*self.release_pages(), now=NOW)
        result = parse_japan_labour(self.book(), release, now=NOW)
        self.assertEqual((result['value'], result['reference_period']), (2.4, '2026-07'))
        self.assertIsNone(result['published_at'])
        self.assertEqual(result['release_date_known'], '2026-08-28')
        self.assertEqual(result['next_due_at'], '2026-10-01T15:00:00+00:00')

    def test_wrong_sex_adjustment_year_and_missing_latest_rejected(self):
        from official_quarterly_labour import parse_japan_labour, parse_japan_labour_release
        release = parse_japan_labour_release(*self.release_pages(), now=NOW)
        mutations = [lambda s: setattr(s.cell(9,20), 'value', 'Male'),
                     lambda s: setattr(s.cell(5,5), 'value', 'Original series'),
                     lambda s: setattr(s.cell(12,1), 'value', 2025),
                     lambda s: setattr(s.cell(17,20), 'value', None),
                     lambda s: setattr(s.cell(18,20), 'value', 2.3),
                     lambda s: setattr(s.cell(17,20), 'value', True)]
        for mutation in mutations:
            with self.subTest(mutation=mutation), self.assertRaises(ValueError):
                parse_japan_labour(self.book(mutation), release, now=NOW)

    def test_release_calendar_mismatch_and_due_release_fail(self):
        from official_quarterly_labour import parse_japan_labour_release
        results, schedule = self.release_pages()
        with self.assertRaisesRegex(ValueError, 'CONFLICT'):
            parse_japan_labour_release(results, schedule.replace('August 28', 'August 29'), now=NOW)
        with self.assertRaisesRegex(ValueError, 'NEW_RELEASE_DUE'):
            parse_japan_labour_release(results, schedule, now=datetime(2026,10,1,15,tzinfo=timezone.utc))
        with self.assertRaisesRegex(ValueError, 'FUTURE_RELEASE'):
            parse_japan_labour_release(results, schedule, now=datetime(2026,8,27,tzinfo=timezone.utc))

    def test_fetch_rejects_failed_calendar_before_downloading_workbook(self):
        from official_quarterly_labour import fetch_japan_labour
        from unittest.mock import Mock
        import requests
        session = Mock(); results, calendar_page = self.release_pages()
        first = Mock(text=results); second = Mock(text=calendar_page)
        second.raise_for_status.side_effect = requests.HTTPError()
        session.get.side_effect = [first, second]
        with self.assertRaises(requests.HTTPError): fetch_japan_labour(now=NOW, session=session)
        self.assertEqual(session.get.call_count, 2)


if __name__ == '__main__': unittest.main()
